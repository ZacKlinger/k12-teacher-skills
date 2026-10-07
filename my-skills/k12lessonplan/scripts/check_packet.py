#!/usr/bin/env python3
"""Check a rendered student packet the way a copier and a student will meet it.

Usage:
    python3 check_packet.py packet.docx [--max-pages 2] [--sheet pages.png]
    python3 check_packet.py multi-day-packet.docx --days 2 [--sheet pages.png]
    python3 check_packet.py exported-from-google-docs.pdf

The packet is printed from Google Docs ("Add to Drive"), and Docs lays the same file out
taller than Word or LibreOffice: measured on these packets, about 4% taller text plus a few
points at every task, so a page that is 96% full here is past 100% there and its last task
moves to the next page. The report estimates the Docs fill for every page and budgets
against that. Given a PDF exported from the Google Doc itself, it measures that directly.

Converts the packet to PDF with LibreOffice, then measures every page: how far down
the ink reaches, and so how much of the page is paper nobody writes on. It reports
the page count against the budget, any page that ends early because a block too tall
to fit was pushed to the next one, and a last page that is mostly empty. With
--sheet it also writes every page side by side into one image, so the whole packet
can be looked at once instead of page by page.

Needs LibreOffice (soffice), poppler (pdftoppm), and Pillow. If any is missing it says
so and exits 0: the check is worth having, not worth blocking a lesson on.

Every task is checked whole. The renderer prints each task (its heading, prompt, language
line, starters, lines, and the table or organizer it is answered in) as one table row, and
Google Docs never splits a row unless it is taller than a page. So the check tags the top
and bottom of every task, finds both on the printed pages, and reports an error for any task
that breaks across a page here, or is too tall to stay whole in Google Docs. A prompt is
never on one page and its answer space on the next.

With --days N (a multi-day packet, only built when one is asked for) it finds where each
day starts and checks each day on its own: every day starts on the front of a fresh sheet,
so a day can be handed out, collected, or reprinted from the Google Doc by itself, and each
day aims for two pages and never passes four. Days are found by the heading that follows
each page break, so it needs pdftotext (poppler) and the .docx, not a PDF.
"""

import argparse
import glob
import html
import os
import re
import shutil
import subprocess
import tempfile
import zipfile

# The renderer's margins are 0.7 in, and the footer sits just above the bottom edge of
# the body, so the measurement stops short of it.
MARGIN_IN = 0.7
FOOTER_IN = 1.0
# Google Docs vs LibreOffice on the same packet: ~4% taller text, ~3pt more per task unit.
DOCS_FACTOR = 1.06
DPI = 40


def ink_bottom(img, top_px, bottom_px):
    """Lowest row with ink between the margins, as a fraction of the writable height."""
    g = img.convert("L")
    w, _ = g.size
    px = g.load()
    for y in range(bottom_px - 1, top_px, -1):
        dark = 0
        for x in range(0, w, 2):
            if px[x, y] < 200:
                dark += 1
                if dark > 2:
                    return (y - top_px) / float(bottom_px - top_px)
    return 0.0


def day_markers(docx_path):
    """The first words after every page break: the heading that opens each later day."""
    xml = zipfile.ZipFile(docx_path).read("word/document.xml").decode("utf-8")
    marks = []
    for chunk in re.split(r'<w:br\b[^>]*w:type="page"[^>]*/>', xml)[1:]:
        words = " ".join(re.findall(r"<w:t(?:\s[^>]*)?>([^<]*)</w:t>", chunk[:40000]))
        marks.append(" ".join(html.unescape(words).split())[:24])
    return marks


def day_starts(pdf, marks):
    """0-based page index where each day starts, found by its opening heading; None for a
    day whose heading can't be found in the text."""
    out = subprocess.run(["pdftotext", "-layout", pdf, "-"], capture_output=True, timeout=60)
    pages = [" ".join(t.split()) for t in out.stdout.decode("utf-8", "replace").split("\f")]
    starts, after = [0], 0
    for m in marks:
        hit = next((i for i in range(after + 1, len(pages)) if m and m in pages[i]), None)
        starts.append(hit)
        after = hit if hit is not None else after
    return starts


MARK = "QZQ{}{}QZQ"     # QZQS12QZQ opens task 12, QZQE12QZQ closes it


def mark_tasks(docx_path, out_path):
    """Copy the packet with an invisible tag (white, 1 pt) at the top and bottom of every
    one-row table: every task unit, the name line, the objective. Returns a label per tag
    number, or None when python-docx isn't installed."""
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
        from docx.table import Table
    except ImportError:
        return None
    doc = Document(docx_path)
    labels = {}
    n = 0
    for el in doc.element.body.iterchildren():
        if not el.tag.endswith("}tbl"):
            continue
        t = Table(el, doc)
        if len(t.rows) != 1:
            continue
        n += 1
        text = " ".join("".join(x.text or "" for x in el.iter() if x.tag.endswith("}t")).split())
        q = re.search(r"(?:^|\s)(\d+[a-z]?)\.\s", text)
        labels[n] = f"the task with question {q.group(1)}" if q else f'"{text[:40]}"'
        first = t.rows[0].cells[0].paragraphs[0]
        last = t.rows[0].cells[-1].paragraphs[-1]
        for par, kind, at_start in ((first, "S", True), (last, "E", False)):
            r = par.add_run(MARK.format(kind, n))
            r.font.size = Pt(1)
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
            if at_start:
                par._p.remove(r._r)
                ppr = par._p.pPr
                (ppr.addnext if ppr is not None else lambda e: par._p.insert(0, e))(r._r)
    doc.save(out_path)
    return labels


def task_breaks(pdf, labels, page_h_pt, margin_in):
    """Errors for every tagged task that crosses a page, or that would not fit on one page
    in Google Docs. Empty when pdftotext is missing."""
    if not labels or not shutil.which("pdftotext"):
        return [], 0
    out = subprocess.run(["pdftotext", "-bbox", pdf, "-"], capture_output=True, timeout=60)
    found = {}
    for page, body in enumerate(out.stdout.decode("utf-8", "replace").split("<page ")[1:]):
        for y0, y1, word in re.findall(
                r'<word xMin="[\d.]+" yMin="([\d.]+)" xMax="[\d.]+" yMax="([\d.]+)">([^<]*)</word>',
                body):
            for m in re.finditer(r"QZQ([SE])(\d+)QZQ", word):
                found.setdefault((m.group(1), int(m.group(2))), (page, float(y0), float(y1)))
    usable = page_h_pt - 2 * margin_in * 72
    errors, checked = [], 0
    for n, label in labels.items():
        a, b = found.get(("S", n)), found.get(("E", n))
        if not a or not b:
            continue
        checked += 1
        if a[0] != b[0]:
            errors.append(f"{label[0].upper() + label[1:]} breaks across pages {a[0] + 1} and "
                          f"{b[0] + 1}: it is taller than a page. Split it into two tasks, or "
                          "shorten its table or organizer, so the prompt and its answer space "
                          "print together.")
        elif (b[2] - a[1]) * DOCS_FACTOR > usable:
            errors.append(f"{label[0].upper() + label[1:]} fills about "
                          f"{(b[2] - a[1]) * DOCS_FACTOR / usable:.0%} of a page in Google "
                          "Docs, so Docs will break it. Split it into two tasks or shorten it.")
    return errors, checked


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("docx")
    ap.add_argument("--max-pages", type=int, default=None,
                    help="page budget: 2 (one sheet, both sides) for a 60-minute lesson")
    ap.add_argument("--days", type=int, default=1,
                    help="a multi-day packet: each day is checked as its own sheet")
    ap.add_argument("--sheet", help="write all pages side by side into this PNG")
    args = ap.parse_args()

    from_docs = args.docx.lower().endswith(".pdf")
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if (not soffice and not from_docs) or not shutil.which("pdftoppm"):
        print("  skip  LibreOffice or pdftoppm is not installed; look at the pages by hand.")
        return 0
    try:
        from PIL import Image
    except ImportError:
        print("  skip  Pillow is not installed; look at the pages by hand.")
        return 0

    work = tempfile.mkdtemp()
    labels = None
    if from_docs:
        pdfs = [os.path.abspath(args.docx)]
    else:
        # the tags are white and 1 pt, so the copy measures the same as the packet
        source = os.path.abspath(args.docx)
        marked = os.path.join(work, "packet.docx")
        try:
            labels = mark_tasks(source, marked)
        except Exception:
            labels = None
        env = dict(os.environ, HOME=os.environ.get("HOME") or work)
        subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", work,
                        marked if labels is not None else source],
                       capture_output=True, env=env, timeout=180)
        pdfs = glob.glob(os.path.join(work, "*.pdf"))
    if not pdfs:
        print("  skip  LibreOffice could not convert the packet; look at the pages by hand.")
        return 0
    subprocess.run(["pdftoppm", "-r", str(DPI), "-png", pdfs[0], os.path.join(work, "p")],
                   capture_output=True, timeout=120)
    pages = [Image.open(f) for f in sorted(glob.glob(os.path.join(work, "p-*.png")))]
    n = len(pages)

    errors, warns = [], []
    fills = []
    if not from_docs:
        breaks, checked = task_breaks(pdfs[0], labels, pages[0].size[1] / DPI * 72, MARGIN_IN)
        if checked:
            print(f"  ok    {checked} task(s) checked whole: no prompt is apart from its "
                  "answer space" if not breaks else f"  ok    {checked} task(s) checked")
            errors.extend(breaks)
        else:
            warns.append("Could not check that every task stays whole (needs python-docx and "
                         "pdftotext); look at the page image for a prompt cut from its lines.")
    day_last = set()       # the last page of every day but the final one ends early on purpose
    if args.days > 1:
        if from_docs or not shutil.which("pdftotext"):
            warns.append("Per-day checks need the .docx and pdftotext; look at where each day "
                         "starts in the page image by hand. Every day starts on an odd page.")
        else:
            marks = day_markers(args.docx)
            if len(marks) != args.days - 1:
                errors.append(f"--days {args.days}, but the packet has {len(marks) + 1} part(s) "
                              "between page breaks. Put one page_break before every day after "
                              "the first, and nowhere else.")
            else:
                starts = day_starts(pdfs[0], marks)
                if None in starts:
                    warns.append("Could not find where every day starts; check the page image: "
                                 "every day starts on an odd page.")
                else:
                    ends = starts[1:] + [n]
                    spans = []
                    for d, (a, b) in enumerate(zip(starts, ends), 1):
                        spans.append(f"day {d} p{a + 1}-{b}")
                        if b - 1 < n - 1:
                            day_last.add(b - 1)
                        if a % 2 == 1:
                            errors.append(
                                f"Day {d} starts on page {a + 1}, the back of a sheet, because day "
                                f"{d - 1} runs {a - starts[d - 2]} pages. Bring day {d - 1} to two "
                                "pages so every day starts on a fresh sheet.")
                        if b - a > 4:
                            errors.append(f"Day {d} is {b - a} pages; a day never passes four.")
                        elif b - a > 2:
                            warns.append(f"Day {d} is {b - a} pages; a day aims for two.")
                    print("  ok    " + ", ".join(spans))
    for img in pages:
        top = int(MARGIN_IN * DPI)
        bottom = img.size[1] - int(FOOTER_IN * DPI)
        fills.append(ink_bottom(img, top, bottom))

    print(f"  ok    {n} page(s): " + ", ".join(f"p{i + 1} {f:.0%}" for i, f in enumerate(fills))
          + ("  (measured on the Google Docs export)" if from_docs else ""))
    if not from_docs:
        docs = [f * DOCS_FACTOR for f in fills]
        print("  ok    in Google Docs, about: "
              + ", ".join(f"p{i + 1} {d:.0%}" for i, d in enumerate(docs)))
        over = [i for i, d in enumerate(docs) if d > 1.0]
        for i in over:
            if i in day_last:
                errors.append(
                    f"Page {i + 1} ends a day and is about {docs[i]:.0%} full in Google Docs: its "
                    "last task will spill onto the next page there and push the next day onto "
                    "the back of a sheet. Free a line or two on it.")
            elif i == n - 1:
                n += 1
                errors.append(
                    f"The last page is {fills[i]:.0%} full here, about {docs[i]:.0%} in Google Docs: "
                    "its last task will print on a page of its own. Free about "
                    f"{int((docs[i] - 0.97) * 30) + 1} line(s) on it.")
            else:
                warns.append(
                    f"Page {i + 1} is {fills[i]:.0%} full here, about {docs[i]:.0%} in Google Docs: "
                    "its last task will likely move to the next page there. Leave a line or two "
                    "of headroom.")
    budget = args.max_pages or (4 * args.days if args.days > 1 else 2)
    if n > budget:
        spare = sum(1 - f for f in fills[:-1]) + (1 - fills[-1])
        errors.append(
            f"{n} pages against a budget of {budget}. The pages hold about "
            f"{spare:.1f} pages of unused paper between them; recover that first (below), "
            "then cut: a section with no student task, a table that repeats another's rows, "
            "a second writing line nobody will fill.")
    for i, f in enumerate(fills[:-1]):
        if f < 0.75 and i not in day_last:
            warns.append(
                f"Page {i + 1} ends at {f:.0%}: the next block was too tall to fit and moved "
                "whole. Move a short block (a note, a word bank, a one-line question) after "
                "it so the gap fills, or shorten the tall block.")
    if n > 1 and fills[-1] < 0.4:
        warns.append(
            f"The last page is {fills[-1]:.0%} full. That is a sheet of paper for a few "
            "lines: pull them back onto the page before, or cut.")
    if n % 2 == 1 and n > 1:
        warns.append(f"{n} pages prints with a blank back. An even count uses both sides.")

    if args.sheet:
        w = sum(p.size[0] for p in pages) + 10 * (n - 1)
        h = max(p.size[1] for p in pages)
        sheet = Image.new("RGB", (w, h), "#808080")
        x = 0
        for p in pages:
            sheet.paste(p, (x, 0))
            x += p.size[0] + 10
        sheet.save(args.sheet)
        print(f"  ok    all pages in one image: {args.sheet}")

    for m in warns:
        print(f"  warn  {m}")
    for m in errors:
        print(f"  ERROR {m}")
    shutil.rmtree(work, ignore_errors=True)
    print()
    if errors:
        print(f"{len(errors)} error(s), {len(warns)} warning(s).")
        return 1
    print(f"clean — 0 errors, {len(warns)} warning(s).")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
