#!/usr/bin/env python3
"""Render a lesson JSON into an editable Word document.

Usage:
    python3 render_packet.py packet.json out.docx

One renderer serves both audiences. `"audience": "student"` produces the student packet
(large type, generous write space, hints, stems). `"audience": "teacher"` produces the
lesson plan (tighter type, phase headers with minutes, notes and tables).

The renderer's main job beyond typesetting is keeping a task whole: a question, its hint,
its stems, and its answer space are bound together so a page break can never land inside
them. Nothing in the JSON needs to say that — it happens for every task block.

Any block a student has to act on may carry an `"es"` string: one short Spanish line,
rendered under the English it supports. It is abbreviated support, not a translation --
see references/packet.md for what gets one and what deliberately does not. On a student
packet the renderer reports, on stderr, every question that has no Spanish line and every
Spanish line that ran longer than its English.

A room with more than one home language lists them in `meta.languages` (default ["es"]),
and each block carries one short line per language under a key named by its code:
`"es": "...", "zh": "..."`. The report then checks every listed language, not only Spanish.

`meta.code` is the teacher's own name for the session ("Science 1.7") and leads the header
and footer. `meta.large_print: true` sets the whole packet in larger type for the students
whose plans call for it.

Requires python-docx. Install with: pip install python-docx --break-system-packages
"""

import json
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

# ---------------------------------------------------------------- style constants

STUDENT = {
    "font": "Verdana",
    "body": 12,
    "h1": 17,
    "h2": 13.5,
    "small": 10,
    "es": 11,             # the language line is read by the students with the least margin
    "margin": 0.7,
    "line_gap": 30,       # points between writing lines — sized for teen handwriting
    "space_after": 8,
}

TEACHER = {
    "font": "Calibri",
    "body": 11,
    "h1": 16,
    "h2": 12.5,
    "small": 9.5,
    "es": 10,
    "margin": 0.75,
    "line_gap": 24,
    "space_after": 6,
}

INK = RGBColor(0x1B, 0x1F, 0x24)
MUTED = RGBColor(0x69, 0x70, 0x79)
# The Spanish line is secondary to the English but it is not fine print: the students
# reading it are the ones with the least margin. Darker than MUTED, lighter than INK.
ES_INK = RGBColor(0x3A, 0x3F, 0x45)
ACCENT = RGBColor(0x3E, 0x6D, 0xA8)
RULE = "E4E7EB"
FILL = "F3F6FA"
BOX_LINE = "9AA2AC"
# Structure is drawn in hairlines, never in fills: a grey fill costs toner on every copy,
# prints as mud on a tired copier, and lowers the contrast of whatever sits on it.
GRID = "8A9099"


# ---------------------------------------------------------------- low-level helpers

def _shade(cell, hexfill):
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:fill"), hexfill)
    cell._tc.get_or_add_tcPr().append(el)


def _cell_borders(cell, color=BOX_LINE, sz=6, sides=("top", "left", "bottom", "right"),
                  style="single"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}")
        if side in sides:
            el.set(qn("w:val"), style)
            el.set(qn("w:sz"), str(sz))
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        borders.append(el)


def _para_border(par, side="bottom", color=BOX_LINE, sz=6):
    pPr = par._p.get_or_add_pPr()
    bdr = pPr.find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        pPr.append(bdr)
    el = OxmlElement(f"w:{side}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), "1")
    el.set(qn("w:color"), color)
    bdr.append(el)


def _no_split(table):
    """Rows stay whole, and the table stays on one page: every paragraph in every row
    but the last is kept with the next, so a break can only land before or after it.
    A vote tally split across two pages is two half-organizers."""
    rows = list(table.rows)
    for i, row in enumerate(rows):
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit")
        trPr.append(el)
        if i < len(rows) - 1:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.keep_with_next = True


def _grid(table, color=GRID, sz=4, header=False):
    """Hairline grid on every cell; a firmer rule under the header row."""
    for r, row in enumerate(table.rows):
        for cell in row.cells:
            _cell_borders(cell, color, sz)
            if header and r == 0:
                tcb = cell._tc.get_or_add_tcPr().find(qn("w:tcBorders"))
                bottom = tcb.find(qn("w:bottom"))
                bottom.set(qn("w:sz"), "10")
                bottom.set(qn("w:color"), "1B1F24")


def _widths(table, inches):
    """Fix column widths. Word ignores a cell width unless autofit is off and every
    cell in the column agrees."""
    table.autofit = False
    for i, w in enumerate(inches):
        table.columns[i].width = Inches(w)
        for cell in table.columns[i].cells:
            cell.width = Inches(w)


def _valign_bottom(cell):
    el = OxmlElement("w:vAlign")
    el.set(qn("w:val"), "bottom")
    cell._tc.get_or_add_tcPr().append(el)


BLANK = re.compile(r"_{2,}")
# a gap wide enough to write a word or two in by hand
GAP = "\u00a0" * 18


def _row_height(row, inches):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(inches * 1440)))
    el.set(qn("w:hRule"), "atLeast")
    trPr.append(el)


# languages written without spaces between words: the word-count length check skips them
UNSPACED = {"zh", "ja", "ko", "th", "my", "km", "lo"}
# languages written right to left: their line is right-aligned and marked bidi
RTL = {"ar", "fa", "ur", "he", "ps", "prs"}


def _rtl(par, lang):
    if str(lang).split("-")[0] not in RTL:
        return
    par.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ppr = par._p.get_or_add_pPr()
    # schema order: bidi sits ahead of spacing, indentation and justification
    ppr.insert_element_before(
        OxmlElement("w:bidi"), "w:adjustRightInd", "w:snapToGrid", "w:spacing", "w:ind",
        "w:contextualSpacing", "w:mirrorIndents", "w:suppressOverlap", "w:jc",
        "w:textDirection", "w:textAlignment", "w:textboxTightWrap", "w:outlineLvl",
        "w:divId", "w:cnfStyle", "w:rPr", "w:sectPr", "w:pPrChange")
    for r in par.runs:
        r._r.get_or_add_rPr().append(OxmlElement("w:rtl"))


class Renderer:
    def __init__(self, data, out_path):
        self.data = data
        self.out = out_path
        self.audience = data.get("audience", "student")
        self.S = dict(STUDENT if self.audience == "student" else TEACHER)
        meta = data.get("meta", {})
        if meta.get("large_print"):
            for k in ("body", "h1", "h2", "small", "es"):
                self.S[k] = round(self.S[k] * 1.25 * 2) / 2
            self.S["line_gap"] = round(self.S["line_gap"] * 1.2)
        # the home languages in the room; each carries one short line under the English
        self.langs = [str(c) for c in (meta.get("languages") or ["es"])]
        self.days_seen = 0      # `day` blocks so far; every one after the first starts a page
        self.doc = Document()
        self._setup()
        # every paragraph the renderer emits is registered here so a task group can be
        # bound together after the fact
        self._group = None
        self._last_group = []
        # where paragraphs and tables go: the document, or the cell of an open unit
        self.target = self.doc
        self._fresh = None
        # language-line coverage, reported on stderr once the document is written
        self._es = []
        self._missing_es = {c: [] for c in self.langs}
        self._long_es = []

    # ------------------------------------------------------------ document setup
    def _setup(self):
        sec = self.doc.sections[0]
        m = Inches(self.S["margin"])
        sec.top_margin = sec.bottom_margin = m
        sec.left_margin = sec.right_margin = m
        normal = self.doc.styles["Normal"]
        normal.font.name = self.S["font"]
        normal.font.size = Pt(self.S["body"])
        normal.font.color.rgb = INK
        rpr = normal.element.get_or_add_rPr()
        rfonts = rpr.get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), self.S["font"])
        pf = normal.paragraph_format
        pf.space_after = Pt(self.S["space_after"])
        pf.space_before = Pt(0)
        pf.line_spacing = 1.15
        self._footer()

    def _footer(self):
        meta = self.data.get("meta", {})
        foot = self.doc.sections[0].footer.paragraphs[0]
        bits = [b for b in (meta.get("code"), meta.get("course"), meta.get("day"),
                            meta.get("title")) if b]
        run = foot.add_run("  ·  ".join(bits))
        run.font.size = Pt(self.S["small"] - 1)
        run.font.color.rgb = MUTED
        run.font.name = self.S["font"]
        foot.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------------------------------------------------ text primitives
    def _par(self):
        """A new paragraph in the current target. A fresh unit cell already holds one
        empty paragraph; the first paragraph written there reuses it, so a unit never
        opens with a blank line."""
        if self._fresh is not None:
            par, self._fresh = self._fresh, None
            return par
        return self.target.add_paragraph()

    def _table(self, rows, cols):
        self._fresh = None
        t = self.target.add_table(rows=rows, cols=cols)
        if self.target is not self.doc:
            # python-docx follows a table in a cell with an empty, full-height
            # paragraph; hand it to whatever is written next instead of leaving a
            # blank line after every table in every unit
            self._fresh = self.target.paragraphs[-1]
        return t

    def p(self, text="", size=None, bold=False, italic=False, color=None,
          space_after=None, space_before=None, indent=0, keep=False, align=None):
        par = self._par()
        if indent:
            par.paragraph_format.left_indent = Inches(indent)
        if space_after is not None:
            par.paragraph_format.space_after = Pt(space_after)
        if space_before is not None:
            par.paragraph_format.space_before = Pt(space_before)
        if align is not None:
            par.alignment = align
        if keep:
            par.paragraph_format.keep_with_next = True
        if text:
            self.run(par, text, size=size, bold=bold, italic=italic, color=color)
        if self._group is not None:
            self._group.append(par)
        return par

    def gap(self, pts=6):
        """The paragraph Word needs after every table, held to a few points. At full
        line height it is a wasted line after every table, and when a table fills a
        page it is the reason a blank page prints."""
        par = self._par()
        pf = par.paragraph_format
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(max(1, int(pts)))
        if self._group is not None:
            self._group.append(par)
        return par

    def run(self, par, text, size=None, bold=False, italic=False, color=None):
        r = par.add_run(text)
        r.font.name = self.S["font"]
        r.font.size = Pt(size or self.S["body"])
        r.bold = bold
        r.italic = italic
        r.font.color.rgb = color or INK
        return r

    def es_line(self, text, indent=0, space_after=4, lang="es"):
        """The Spanish support line — one line, under the English it supports.

        Not italic. Italic costs a striving reader real speed, and this line is read
        by the students who have the least of it to spare; the separation from the
        English comes from size and colour instead.
        """
        if not text:
            return None
        self._es.append(str(text))
        par = self.p(space_after=space_after, space_before=0, indent=indent)
        self.rich(par, str(text), size=self.S["es"])
        for r in par.runs:
            r.font.color.rgb = ES_INK
        _rtl(par, lang)
        return par

    def es_in_cell(self, cell, text, lang="es"):
        """Same line, but inside a tinted box (a note, a word bank) so it stays in it."""
        if not text:
            return None
        self._es.append(str(text))
        par = cell.add_paragraph()
        par.paragraph_format.space_before = Pt(2)
        par.paragraph_format.space_after = Pt(0)
        self.rich(par, str(text), size=self.S["es"])
        for r in par.runs:
            r.font.color.rgb = ES_INK
        _rtl(par, lang)
        return par

    def lang_lines(self, blk, indent=0, space_after=4):
        """One line per home language, in the order meta.languages lists them."""
        pars = []
        for code in self.langs:
            par = self.es_line(blk.get(code), indent=indent, space_after=space_after, lang=code)
            if par is not None:
                pars.append(par)
        return pars

    def lang_in_cell(self, cell, blk):
        for code in self.langs:
            self.es_in_cell(cell, blk.get(code), lang=code)

    def rich(self, par, text, size=None):
        """Renders **bold** spans inside a plain string. Everything else is literal."""
        for i, chunk in enumerate(text.split("**")):
            if chunk:
                self.run(par, chunk, size=size, bold=(i % 2 == 1))

    # ------------------------------------------------------------ group binding
    def begin_group(self):
        self._group = []

    def end_group(self):
        """Bind everything emitted since begin_group so a page break can't split it."""
        self._last_group = list(self._group or [])
        if not self._group:
            self._group = None
            return
        for par in self._group[:-1]:
            par.paragraph_format.keep_with_next = True
        self._group = None

    # ------------------------------------------------------------ blocks
    def heading(self, text, minutes=None, es=None, blk=None):
        """One line per section: the name, its gloss in the room's language, and the
        minutes flush right. Three lines of heading per section is a page of headings
        across a packet."""
        par = self.p(space_before=10, space_after=4, keep=True)
        self.run(par, text.upper(), size=self.S["h2"], bold=True)
        lines = blk if blk is not None else {"es": es}
        # with one home language the gloss rides on the heading line; with two or
        # more, headings go without, and the lines stay on the tasks (SKILL.md)
        if len(self.langs) == 1 and lines.get(self.langs[0]):
            gloss = self.run(par, "  ·  " + str(lines[self.langs[0]]), size=self.S["es"])
            gloss.font.color.rgb = ES_INK
        if minutes:
            width = 8.5 - 2 * self.S["margin"]
            par.paragraph_format.tab_stops.add_tab_stop(Inches(width), WD_TAB_ALIGNMENT.RIGHT)
            self.run(par, f"\t{minutes} min", size=self.S["small"], color=MUTED, bold=True)
        _para_border(par, "bottom", GRID, 6)
        par.paragraph_format.keep_with_next = True
        return par

    def title_block(self):
        meta = self.data.get("meta", {})
        eyebrow = "  ·  ".join(
            b for b in (meta.get("code"), meta.get("course"), meta.get("day"), meta.get("period"))
            if b
        )
        if eyebrow:
            par = self.p(space_after=2, keep=True)
            self.run(par, eyebrow.upper(), size=self.S["small"], color=MUTED, bold=True)
        par = self.p(space_after=6, keep=True)
        self.run(par, meta.get("title", ""), size=self.S["h1"], bold=True)

        if self.audience == "student" and meta.get("name_line", True):
            self.name_line()

        obj = self.data.get("objective")
        if obj:
            self._banner("Objective", obj, self.data.get("standard"))
        # The agenda belongs on the board and in the deck, not on the student's page —
        # printing it there costs a third of page 1 and students read it off the screen
        # anyway. The lesson plan always prints it.
        agenda = self.data.get("agenda")
        if agenda and (self.audience == "teacher" or meta.get("show_agenda")):
            self.agenda(agenda)

    def name_line(self):
        t = self._table(rows=1, cols=2)
        t.autofit = True
        for cell, label in zip(t.rows[0].cells, ("Name:", "Date:")):
            cell.text = ""
            par = cell.paragraphs[0]
            self.run(par, label, size=self.S["small"], color=MUTED)
            par.paragraph_format.space_after = Pt(2)
            _cell_borders(cell, RULE, 6, sides=("bottom",))
        _no_split(t)
        self.gap(4)

    def day(self, blk):
        """Opens one day of a multi-day packet. Every day after the first starts on a fresh
        page with its own name line, so a day can be handed out, collected, or reprinted from
        the Google Doc on its own; each day carries its own "I can"."""
        later = self.days_seen > 0
        self.days_seen += 1
        if later:
            self.page_break()
        eyebrow = "  ·  ".join(str(b) for b in (blk.get("code"), blk.get("day"), blk.get("period"))
                               if b)
        if eyebrow:
            par = self.p(space_before=0 if later else 6, space_after=2, keep=True)
            self.run(par, eyebrow.upper(), size=self.S["small"], color=MUTED, bold=True)
        if blk.get("title"):
            par = self.p(space_after=4, keep=True)
            self.run(par, blk["title"], size=self.S["h2"] + 1.5, bold=True)
            line = blk.get(self.langs[0]) if len(self.langs) == 1 else None
            if line:
                gloss = self.run(par, "  ·  " + str(line), size=self.S["es"])
                gloss.font.color.rgb = ES_INK
        meta = self.data.get("meta", {})
        if later and self.audience == "student" and meta.get("name_line", True):
            self.name_line()
        if blk.get("objective"):
            self._banner("Objective", blk["objective"], blk.get("standard"))

    def page_break(self):
        # The break's own paragraph mark lands at the top of the new page, so it is made
        # 1 pt tall instead of a blank line.
        par = self._par()
        par.add_run().add_break(WD_BREAK.PAGE)
        pf = par.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(1)
        pf.space_before = pf.space_after = Pt(0)

    def _banner(self, label, text, sub=None):
        t = self._table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        _cell_borders(cell, "1B1F24", 12, sides=("left",))
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        self.run(par, ("I can" if label.lower() == "objective" else label).upper() + "  ",
                 size=self.S["small"], bold=True)
        body = text[6:] if label.lower() == "objective" and text.lower().startswith("i can ") else text
        self.run(par, body, size=self.S["body"])
        if sub:
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(0)
            self.run(p2, sub, size=self.S["small"], color=MUTED)
        _no_split(t)
        self.gap(2)

    def agenda(self, items):
        par = self.p(space_before=8, space_after=3, keep=True)
        self.run(par, "TODAY", size=self.S["small"], bold=True, color=MUTED)
        rows = []
        for it in items:
            if isinstance(it, (list, tuple)):
                rows.append((str(it[0]), f"{it[1]} min" if len(it) > 1 else ""))
            else:
                rows.append((str(it), ""))
        t = self._table(rows=len(rows), cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        for (name, mins), row in zip(rows, t.rows):
            row.cells[0].width = Inches(4.6)
            row.cells[1].width = Inches(1.1)
            p1 = row.cells[0].paragraphs[0]
            p1.paragraph_format.space_after = Pt(1)
            self.run(p1, name, size=self.S["body"])
            p2 = row.cells[1].paragraphs[0]
            p2.paragraph_format.space_after = Pt(1)
            p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            self.run(p2, mins, size=self.S["small"], color=MUTED)
            for c in row.cells:
                _cell_borders(c, RULE, 4, sides=("bottom",))
        _no_split(t)
        self.gap(4)

    def question(self, blk):
        self.begin_group()
        num = str(blk.get("number", "")).strip()
        par = self.p(space_before=9, space_after=3)
        if num:
            self.run(par, f"{num}. ", bold=True)
        self.rich(par, blk.get("prompt", ""))
        if blk.get("minutes"):
            self.run(par, f"   ({blk['minutes']} min)", size=self.S["small"], color=MUTED)

        for code in self.langs:
            line = blk.get(code)
            if line:
                self.es_line(line, lang=code)
                self._check_es_length(blk.get("prompt", ""), line, num, code)
            elif self.audience == "student":
                self._missing_es[code].append(num or blk.get("prompt", "")[:40])

        if blk.get("example"):
            self.tinted("Example", blk["example"])
        if blk.get("hint"):
            self.tinted("Hint", blk["hint"])
        for sub in blk.get("parts", []) or []:
            par = self.p(space_after=3, indent=0.3)
            self.rich(par, sub)
        choices = blk.get("choices") or []
        if choices:
            # options to circle get a line of their own, spaced wide enough to
            # circle one without touching the next
            par = self.p(space_before=4, space_after=4, indent=0.3)
            for i, c in enumerate(choices):
                if i:
                    self.run(par, "\u00a0" * 10)
                self.run(par, str(c), bold=True)
        stems = blk.get("stems", []) or []
        spec = blk.get("space", {"kind": "lines", "count": 3})
        if stems and isinstance(spec, dict) and spec.get("kind", "lines") == "lines":
            # the stem is the first writing line, not a box above the lines: the
            # student starts writing where the sentence starts
            self.space(dict(spec, count=max(int(spec.get("count", 2)), len(stems))), stems)
        else:
            for stem in stems:
                self.stem(stem)
            self.space(spec)
        if spec == {"kind": "none"} or (isinstance(spec, dict) and spec.get("kind") == "none"):
            # a question whose answer space is the table or organizer after it
            # stays on the same page as that table
            for par in self._group or []:
                par.paragraph_format.keep_with_next = True
        self.end_group()

    def stem(self, text):
        t = self._table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        _cell_borders(cell, GRID, 8, sides=("left",), style="single")
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        par.paragraph_format.space_before = Pt(1)
        self.rich(par, text)
        _no_split(t)
        p = self.p(space_after=2)
        return t

    def tinted(self, label, text):
        par = self.p(space_after=4, indent=0.15)
        self.run(par, f"{label}: ", size=self.S["small"], bold=True, color=ACCENT)
        self.rich(par, text, size=self.S["small"])
        return par

    def stem_on_line(self, cell, text):
        """A sentence frame printed on its writing line: blanks become gaps on the
        rule, and a trailing blank becomes the rest of the line."""
        text = BLANK.sub("___", str(text)).strip()
        text = re.sub(r"\s*___\s*[.?!]?$", "", text)   # a trailing blank is the rest of the line
        par = cell.paragraphs[0]
        for i, chunk in enumerate(text.split("___")):
            if i:
                gap = self.run(par, GAP)
                gap.font.underline = True
            if chunk:
                self.rich(par, chunk)
        _valign_bottom(cell)

    def space(self, spec, stems=None):
        if not spec:
            return
        kind = spec.get("kind", "lines") if isinstance(spec, dict) else str(spec)
        if kind == "none":
            return
        if kind == "lines":
            # Writing lines are a borderless table with a rule under each row. Bordered
            # empty paragraphs look right in isolation but Word merges consecutive
            # identical borders into one box, so three lines print as one.
            count = int(spec.get("count", 3)) if isinstance(spec, dict) else 3
            t = self._table(rows=count, cols=1)
            for i, row in enumerate(t.rows):
                cell = row.cells[0]
                _cell_borders(cell, BOX_LINE, 6, sides=("bottom",))
                par = cell.paragraphs[0]
                par.paragraph_format.space_before = Pt(0)
                par.paragraph_format.space_after = Pt(1)
                if stems and i < len(stems):
                    self.stem_on_line(cell, stems[i])
                _row_height(row, self.S["line_gap"] / 72.0)
            _no_split(t)
            self.gap(6)
            return
        # box / work space
        height = float(spec.get("height_in", 2.0)) if isinstance(spec, dict) else 2.0
        label = spec.get("label") if isinstance(spec, dict) else None
        t = self._table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        _cell_borders(cell, BOX_LINE, 6)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        if label:
            self.run(par, label, size=self.S["small"], color=MUTED)
        _row_height(t.rows[0], height)
        _no_split(t)
        self.gap(4)

    def table(self, blk, fill_rows=0, row_height=None):
        headers = blk.get("headers") or []
        rows = [list(r) for r in (blk.get("rows") or [])]
        blanks = int(blk.get("blank_rows", fill_rows) or 0)
        ncols = len(headers) or (len(rows[0]) if rows else 2)
        total = len(rows) + blanks + (1 if headers else 0)
        if total == 0:
            return
        t = self._table(rows=total, cols=ncols)
        _grid(t, header=bool(headers))
        idx = 0
        if headers:
            for c, head in enumerate(headers):
                cell = t.rows[0].cells[c]
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                self.run(par, str(head), size=self.S["small"], bold=True)
            idx = 1
        for r, data_row in enumerate(rows):
            for c in range(ncols):
                cell = t.rows[idx + r].cells[c]
                par = cell.paragraphs[0]
                par.paragraph_format.space_after = Pt(2)
                val = data_row[c] if c < len(data_row) else ""
                self.rich(par, str(val))
        # every row a student writes in is tall enough to write in, filled or not
        h = float(row_height or blk.get("row_height_in", 0.36))
        for r in range(1 if headers else 0, total):
            _row_height(t.rows[r], h)
            for cell in t.rows[r].cells:
                cell.paragraphs[0].paragraph_format.space_before = Pt(3)
        _no_split(t)
        self.gap(6)

    def note(self, blk):
        t = self._table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        _cell_borders(cell, BOX_LINE, 12, sides=("left",))
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        label = blk.get("label")
        if label:
            self.run(par, label.upper() + "  ", size=self.S["small"], bold=True)
        self.rich(par, blk.get("text", ""))
        self.lang_in_cell(cell, blk)
        _no_split(t)
        self.gap(4)

    def listing(self, blk):
        self.begin_group()
        if blk.get("label"):
            par = self.p(space_before=8, space_after=3)
            self.run(par, blk["label"], bold=True)
        self.lang_lines(blk, space_after=5)
        ordered = blk.get("ordered", False)
        for i, item in enumerate(blk.get("items", []), 1):
            par = self.p(space_after=4, indent=0.25)
            marker = f"{i}. " if ordered else "•  "
            self.run(par, marker, bold=ordered)
            self.rich(par, str(item))
        self.end_group()

    def wordbank(self, blk):
        t = self._table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        _cell_borders(cell, GRID, 4)
        par = cell.paragraphs[0]
        par.paragraph_format.space_after = Pt(0)
        self.run(par, (blk.get("label") or "Word bank").upper() + "  ",
                 size=self.S["small"], bold=True)
        self.run(par, "     ".join(str(i) for i in blk.get("items", [])))
        self.lang_in_cell(cell, blk)
        _no_split(t)
        self.gap(4)


    # ------------------------------------------------------------ graphic organizers
    def organizer(self, blk):
        """A graphic organizer matched to the thinking the task asks for: compare
        (tchart), notice and wonder, sequence (flow), argue (cer), a new word (frayer).
        Drawn in hairlines like everything else, and kept on one page."""
        kind = blk.get("kind", "tchart")
        label = blk.get("label") or (f"New word: {blk['word']}" if blk.get("word") else None)
        if label:
            par = self.p(space_before=8, space_after=2, keep=True)
            self.rich(par, "**" + str(label) + "**")
        for es_par in self.lang_lines(blk, space_after=3):
            es_par.paragraph_format.keep_with_next = True
        if kind in ("tchart", "notice_wonder"):
            cols = blk.get("columns") or (["I notice", "I wonder"] if kind == "notice_wonder"
                                         else ["", ""])
            nrows = int(blk.get("rows", 4))
            t = self._table(rows=nrows + 1, cols=len(cols))
            for c, head in enumerate(cols):
                cell = t.rows[0].cells[c]
                _cell_borders(cell, "1B1F24", 10, sides=("bottom",)
                              + (("right",) if c < len(cols) - 1 else ()))
                self.rich(cell.paragraphs[0], "**" + str(head) + "**")
            for r in range(1, nrows + 1):
                _row_height(t.rows[r], self.S["line_gap"] / 72.0)
                for c in range(len(cols)):
                    _cell_borders(t.rows[r].cells[c], BOX_LINE, 6, sides=("bottom",)
                                  + (("right",) if c < len(cols) - 1 else ()))
        elif kind == "flow":
            steps = blk.get("steps") or ["", "", ""]
            t = self._table(rows=1, cols=len(steps) * 2 - 1)
            arrow_w = 0.32
            box_w = (8.5 - 2 * self.S["margin"] - arrow_w * (len(steps) - 1)) / len(steps)
            _widths(t, [box_w if c % 2 == 0 else arrow_w for c in range(len(steps) * 2 - 1)])
            for i, step in enumerate(steps):
                cell = t.rows[0].cells[i * 2]
                _cell_borders(cell, GRID, 4)
                if step:
                    self.run(cell.paragraphs[0], str(step), size=self.S["small"], bold=True)
                if i < len(steps) - 1:
                    arrow = t.rows[0].cells[i * 2 + 1]
                    _cell_borders(arrow, GRID, 4, sides=())
                    ap = arrow.paragraphs[0]
                    ap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    self.run(ap, "\u2192", size=self.S["h2"], color=MUTED)
                    el = OxmlElement("w:vAlign"); el.set(qn("w:val"), "center")
                    arrow._tc.get_or_add_tcPr().append(el)
            _row_height(t.rows[0], float(blk.get("height_in", 1.1)))
        elif kind == "cer":
            parts = [("Claim", "What I think"), ("Evidence", "What I saw or measured"),
                     ("Reasoning", "Why the evidence proves it")]
            stems = blk.get("stems") or {}
            t = self._table(rows=3, cols=2)
            _grid(t)
            _widths(t, [1.45, 8.5 - 2 * self.S["margin"] - 1.45])
            for r, (name, gloss) in enumerate(parts):
                head, body = t.rows[r].cells
                self.run(head.paragraphs[0], name, bold=True)
                gp = head.add_paragraph()
                self.run(gp, gloss, size=self.S["small"], color=MUTED)
                if stems.get(name.lower()):
                    self.stem_on_line(body, stems[name.lower()])
                _row_height(t.rows[r], float(blk.get("height_in", 0.85)))
        elif kind == "frayer":
            word = blk.get("word", "")
            cells = blk.get("cells") or ["What it means", "Draw it", "An example", "Not an example"]
            t = self._table(rows=2, cols=2)
            _grid(t)
            for i, name in enumerate(cells[:4]):
                cell = t.rows[i // 2].cells[i % 2]
                self.run(cell.paragraphs[0], str(name), size=self.S["small"], bold=True, color=MUTED)
            for row in t.rows:
                _row_height(row, float(blk.get("height_in", 1.25)))
            if word:
                # the word sits above the square, where the eye starts
                pass
        else:
            raise ValueError(f"unknown organizer kind: {kind!r}")
        _no_split(t)
        self.gap(6)

    # ------------------------------------------------------------ spanish coverage
    @staticmethod
    def _words(text):
        return len(re.sub(r"\*\*", "", str(text)).split())

    def _check_es_length(self, english, spanish, num, code="es"):
        """Abbreviated means shorter. Spanish runs 15-20% longer than English for the
        same content, so only a real overshoot is worth reporting. Languages written
        without spaces between words can't be counted this way and are left to a human."""
        if code.split("-")[0] in UNSPACED:
            return
        en, es = self._words(english), self._words(spanish)
        if en and es > en * 1.2:
            name = "Spanish" if code == "es" else f"'{code}'"
            self._long_es.append(f"{num or '?'}: {es} {name} words for {en} English")

    def report(self):
        if self.audience != "student":
            return
        for code, missing in self._missing_es.items():
            if not missing:
                continue
            name = "Spanish" if code == "es" else f"'{code}'"
            print(
                f"  note  no {name} line on question(s): "
                + ", ".join(str(m) for m in missing)
                + "\n        Every question a student answers carries one. See packet.md.",
                file=sys.stderr,
            )
        for msg in self._long_es:
            print(
                f"  note  language line longer than the English it supports — {msg}. "
                "Cut it to the task itself.",
                file=sys.stderr,
            )


    # ------------------------------------------------------------ units
    # Google Docs, which is how these packets are printed, ignores "keep with next"
    # and lets a table break anywhere -- tested: a question and its writing lines,
    # built as a paragraph plus a table, split across pages in Docs every time. The
    # one thing Docs will not split is a single table row. So every task travels in
    # its own borderless one-row table: heading, question, language line, starter,
    # lines, and the table or organizer that answers it, as one piece.
    def begin_unit(self):
        t = self.doc.add_table(rows=1, cols=1)
        _widths(t, [8.5 - 2 * self.S["margin"]])
        cell = t.rows[0].cells[0]
        _cell_borders(cell, "FFFFFF", 0, sides=())
        tcPr = cell._tc.get_or_add_tcPr()
        mar = OxmlElement("w:tcMar")
        for side in ("top", "left", "bottom", "right"):
            el = OxmlElement(f"w:{side}")
            el.set(qn("w:w"), "0")
            el.set(qn("w:type"), "dxa")
            mar.append(el)
        tcPr.append(mar)
        trPr = t.rows[0]._tr.get_or_add_trPr()
        trPr.append(OxmlElement("w:cantSplit"))
        self.target = cell
        self._fresh = cell.paragraphs[0]
        return t

    def end_unit(self):
        # a cell must end in a paragraph; if the last thing in it is a table, the
        # trailing one is held to a point so it adds no line
        cell = self.target
        last = cell.paragraphs[-1]
        if not last.text.strip():
            pf = last.paragraph_format
            pf.space_before = Pt(0)
            pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(1)
        self.target = self.doc
        self._fresh = None
        self.gap(4)

    @staticmethod
    def units(sections):
        """Group blocks into units that print as one piece: a heading, any lead-ins
        after it (a direction, a note, a word bank), the task they lead into, and the
        table or organizer that task is answered in."""
        tabular = ("table", "fill_table", "organizer")
        lead = ("heading", "phase", "text", "paragraph", "labeled", "note", "callout",
                "wordbank", "list", "steps", "stem")
        out, unit = [], []
        i = 0
        while i < len(sections):
            b = sections[i]
            kind = b.get("type")
            if kind == "page_break":
                if unit:
                    out.append(unit)
                    unit = []
                out.append([b])
                i += 1
                continue
            if kind in ("heading", "phase") and unit and \
                    any(u.get("type") not in ("heading", "phase") for u in unit):
                out.append(unit)   # a new section closes a unit still waiting for its task
                unit = []
            unit.append(b)
            i += 1
            if kind in lead and i < len(sections):
                continue           # a lead-in waits for the task it introduces
            spec = b.get("space") if kind == "question" else None
            if kind == "question" and isinstance(spec, dict) and spec.get("kind") == "none" \
                    and i < len(sections) and sections[i].get("type") in tabular:
                unit.append(sections[i])
                i += 1
            out.append(unit)
            unit = []
        if unit:
            out.append(unit)
        return out

    # ------------------------------------------------------------ dispatch
    def render(self):
        self.title_block()
        for unit in self.units(self.data.get("sections", [])):
            wrap = self.audience == "student" and any(
                b.get("type") in ("question", "table", "fill_table", "organizer", "heading", "phase")
                for b in unit)
            if wrap:
                self.begin_unit()
            for blk in unit:
                self.block(blk)
            if wrap:
                self.end_unit()
        last = self.doc.paragraphs[-1] if self.doc.paragraphs else None
        if last is not None and not last.text.strip():
            last.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            last.paragraph_format.line_spacing = Pt(1)
            last.paragraph_format.space_after = Pt(0)
            last.paragraph_format.keep_with_next = False
        self.doc.save(self.out)
        self.report()

    def block(self, blk):
        kind = blk.get("type")
        if kind == "heading":
            self.heading(blk.get("text", ""), blk.get("minutes"), blk=blk)
        elif kind == "phase":
            self.heading(blk.get("name", ""), blk.get("minutes"), blk=blk)
        elif kind == "question":
            self.question(blk)
        elif kind in ("text", "paragraph"):
            self.begin_group()
            par = self.p(space_after=6)
            self.rich(par, blk.get("text", ""))
            self.lang_lines(blk, space_after=6)
            self.end_group()
        elif kind == "labeled":
            self.begin_group()
            par = self.p(space_after=6)
            self.run(par, blk.get("label", "") + "  ", bold=True)
            self.rich(par, blk.get("text", ""))
            self.lang_lines(blk, space_after=6)
            self.end_group()
        elif kind in ("list", "steps"):
            blk = dict(blk)
            if kind == "steps":
                blk["ordered"] = True
            self.listing(blk)
        elif kind == "table":
            self.table(blk)
        elif kind == "fill_table":
            self.table(blk)
        elif kind in ("note", "callout"):
            self.note(blk)
        elif kind == "stem":
            self.stem(blk.get("text", ""))
        elif kind == "wordbank":
            self.wordbank(blk)
        elif kind == "organizer":
            self.organizer(blk)
        elif kind == "space":
            self.space(blk)
        elif kind == "page_break":
            self.page_break()
        elif kind == "day":
            self.day(blk)
        elif kind == "spacer":
            self.p(space_after=int(blk.get("points", 12)))
        else:
            raise ValueError(f"unknown block type: {kind!r}")


ACCENT_HEX = "3E6DA8"


def main():
    if len(sys.argv) != 3:
        print(__doc__)
        sys.exit(1)
    with open(sys.argv[1]) as fh:
        data = json.load(fh)
    Renderer(data, sys.argv[2]).render()
    print(f"wrote {sys.argv[2]}")


if __name__ == "__main__":
    main()
